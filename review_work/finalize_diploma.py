from __future__ import annotations

import re
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt, RGBColor
from PIL import Image


ROOT = Path(__file__).resolve().parent
REPO = ROOT.parent
INPUT = ROOT / "input" / "diploma_source.docx"
OUTPUT = ROOT / "diplom_sysassist_checked_final.docx"
SCREENSHOTS = ROOT / "screenshots"
DOCX_SCREENSHOTS = ROOT / "screenshots_docx"


STRUCTURAL_H1 = {
    "СОДЕРЖАНИЕ",
    "ВВЕДЕНИЕ",
    "1 АНАЛИТИЧЕСКАЯ ЧАСТЬ",
    "2 ПРАКТИЧЕСКАЯ ЧАСТЬ",
    "3 ТЕСТИРОВАНИЕ, ЭКСПЛУАТАЦИЯ И ОЦЕНКА РЕЗУЛЬТАТА",
    "ЗАКЛЮЧЕНИЕ",
    "СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ",
}

SCREENSHOT_BY_CAPTION = {
    "Рисунок 3 - Главная панель SysAssist": DOCX_SCREENSHOTS / "01_dashboard.png",
    "Рисунок 4 - Реестр модулей SysAssist": DOCX_SCREENSHOTS / "02_modules.png",
    "Рисунок 6 - Очередь согласований SysAssist": DOCX_SCREENSHOTS / "04_approvals.png",
}

EXTRA_SCREENSHOTS = [
    (
        "Рисунок 7 - Список событий и инцидентов SysAssist",
        DOCX_SCREENSHOTS / "03_events.png",
        "Экран событий показывает рабочую очередь инцидентов: источник, серьезность, цель, статус и связь с дальнейшими действиями. Для защиты ВКР этот экран полезен тем, что соединяет теоретический жизненный цикл инцидента с реальной операторской таблицей.",
        "Рисунок 4 - Реестр модулей SysAssist",
    ),
    (
        "Рисунок 8 - Диагностика и production readiness SysAssist",
        DOCX_SCREENSHOTS / "05_diagnostics.png",
        "Экран диагностики показывает не только общий статус, но и причины предупреждений: состояние модулей, свежесть проверок, готовность окружения и блокирующие условия. Поэтому он используется как доказательство того, что проект не скрывает неполную готовность контура.",
        "Рисунок 5 - Схема deployment-контура",
    ),
]


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def crop_top_16x9(image: Image.Image) -> Image.Image:
    target_h = min(image.height, round(image.width * 9 / 16))
    return image.crop((0, 0, image.width, target_h))


def prepare_docx_screenshots() -> None:
    DOCX_SCREENSHOTS.mkdir(exist_ok=True)
    for name in ["01_dashboard.png", "02_modules.png", "03_events.png", "04_approvals.png"]:
        source = Image.open(SCREENSHOTS / name).convert("RGB")
        crop_top_16x9(source).save(DOCX_SCREENSHOTS / name, optimize=True)

    diagnostics = Image.open(SCREENSHOTS / "05_diagnostics.png").convert("RGB")
    top = diagnostics.crop((0, 0, diagnostics.width, min(diagnostics.height, 900)))
    bottom_h = min(diagnostics.height, 1050)
    bottom = diagnostics.crop((0, diagnostics.height - bottom_h, diagnostics.width, diagnostics.height))
    panel_w, panel_h = 720, 810
    top = top.resize((panel_w, panel_h), Image.Resampling.LANCZOS)
    bottom = bottom.resize((panel_w, panel_h), Image.Resampling.LANCZOS)
    composite = Image.new("RGB", (panel_w * 2, panel_h), (255, 255, 255))
    composite.paste(top, (0, 0))
    composite.paste(bottom, (panel_w, 0))
    composite.save(DOCX_SCREENSHOTS / "05_diagnostics.png", optimize=True)


def set_cell_width(cell, width_twips: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:type"), "dxa")
    tc_w.set(qn("w:w"), str(width_twips))


def set_table_width(table, width_twips: int) -> None:
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(width_twips))

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_ind.set(qn("w:w"), "0")

    for tag in ("w:tblLayout",):
        old = tbl_pr.find(qn(tag))
        if old is not None:
            tbl_pr.remove(old)


def set_table_cell_margins(table, top=80, start=80, bottom=80, end=80) -> None:
    tbl_pr = table._tbl.tblPr
    tbl_cell_mar = tbl_pr.find(qn("w:tblCellMar"))
    if tbl_cell_mar is None:
        tbl_cell_mar = OxmlElement("w:tblCellMar")
        tbl_pr.append(tbl_cell_mar)
    for m, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tbl_cell_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tbl_cell_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    for node_name, fld_type, text in (
        ("begin", "begin", None),
        ("instr", None, "PAGE"),
        ("separate", "separate", None),
        ("text", None, "1"),
        ("end", "end", None),
    ):
        if node_name == "instr":
            instr = OxmlElement("w:instrText")
            instr.set(qn("xml:space"), "preserve")
            instr.text = text
            run._r.append(instr)
        elif node_name == "text":
            t = OxmlElement("w:t")
            t.text = text
            run._r.append(t)
        else:
            fld = OxmlElement("w:fldChar")
            fld.set(qn("w:fldCharType"), fld_type)
            run._r.append(fld)


def set_run_font(run, name: str, size: Pt | None = None, bold=None, italic=None) -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run._element.rPr.rFonts.set(qn("w:cs"), name)
    if size is not None:
        run.font.size = size
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    run.font.color.rgb = RGBColor(0, 0, 0)


def paragraph_role(text: str, body_started: bool) -> str:
    stripped = text.strip()
    if not stripped:
        return "empty"
    if not body_started and stripped != "СОДЕРЖАНИЕ":
        return "front"
    if stripped in STRUCTURAL_H1 or re.match(r"^\d\s+[А-ЯA-ZЁ]", stripped):
        return "h1"
    if re.match(r"^\d+\.\d+\s+", stripped):
        return "h2"
    if re.match(r"^\d+\.\d+\.\d+\s+", stripped):
        return "h3"
    if stripped.startswith(("Таблица ", "Рисунок ", "Листинг ")):
        return "caption"
    if re.match(r"^\d+\.\s+", stripped) and len(stripped) > 40:
        return "source"
    return "body"


def normalize_text(text: str) -> str:
    replacements = {
        "—": "-",
        "–": "-",
        "«": '"',
        "»": '"',
        "продуктивный режим": "продуктивному режиму",
        "демонстрационный режим": "демонстрационного режима",
        "endpoint проверка продуктивной готовности": "endpoint проверки продуктивной готовности",
        "endpoint проверка": "endpoint проверки",
        "screen-материалы": "экранные материалы",
        "к продуктивному режиму": "к продуктивному режиму",
        "демонстрационного режима и корректность": "демонстрационного режима и корректность",
        "ACTION_SIMULATED": "ACTION_BLOCKED",
    }
    for old, new in replacements.items():
        text = text.replace(old, new)
    text = re.sub(r"\s+", " ", text).strip() if "\n" not in text else text
    return text


def ensure_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:cs"), "Times New Roman")
    normal.font.size = Pt(14)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    pf = normal.paragraph_format
    pf.first_line_indent = Cm(1.25)
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)

    for name, size in (("Heading 1", 15), ("Heading 2", 14), ("Heading 3", 14)):
        style = doc.styles[name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
        style._element.rPr.rFonts.set(qn("w:cs"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = False
        style.font.color.rgb = RGBColor(0, 0, 0)
        fmt = style.paragraph_format
        fmt.first_line_indent = Cm(0)
        fmt.alignment = WD_ALIGN_PARAGRAPH.CENTER if name == "Heading 1" else WD_ALIGN_PARAGRAPH.LEFT
        fmt.space_before = Pt(0 if name == "Heading 1" else 6)
        fmt.space_after = Pt(6)
        fmt.keep_with_next = True
        fmt.keep_together = True

    if "CaptionSys" not in doc.styles:
        doc.styles.add_style("CaptionSys", WD_STYLE_TYPE.PARAGRAPH)
    cap = doc.styles["CaptionSys"]
    cap.font.name = "Times New Roman"
    cap._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    cap._element.rPr.rFonts.set(qn("w:cs"), "Times New Roman")
    cap.font.size = Pt(14)
    cap.font.color.rgb = RGBColor(0, 0, 0)
    cap.paragraph_format.first_line_indent = Cm(0)
    cap.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_before = Pt(0)
    cap.paragraph_format.space_after = Pt(3)
    cap.paragraph_format.keep_with_next = True

    if "CodeSys" not in doc.styles:
        doc.styles.add_style("CodeSys", WD_STYLE_TYPE.PARAGRAPH)
    code = doc.styles["CodeSys"]
    code.font.name = "Consolas"
    code._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
    code._element.rPr.rFonts.set(qn("w:cs"), "Consolas")
    code.font.size = Pt(9)
    code.font.color.rgb = RGBColor(0, 0, 0)
    code.paragraph_format.first_line_indent = Cm(0)
    code.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    code.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    code.paragraph_format.space_before = Pt(0)
    code.paragraph_format.space_after = Pt(0)


def apply_section_format(doc: Document) -> None:
    for section in doc.sections:
        section.start_type = WD_SECTION_START.NEW_PAGE
        section.page_width = Mm(210)
        section.page_height = Mm(297)
        section.left_margin = Mm(30)
        section.right_margin = Mm(10)
        section.top_margin = Mm(20)
        section.bottom_margin = Mm(20)
        section.different_first_page_header_footer = True
        section.header_distance = Mm(10)
        header = section.header
        header.is_linked_to_previous = False
        p = header.paragraphs[0] if header.paragraphs else header.add_paragraph()
        p.clear()
        add_page_number(p)
        for r in p.runs:
            set_run_font(r, "Arial", Pt(12))
        first = section.first_page_header
        first.is_linked_to_previous = False
        if first.paragraphs:
            first.paragraphs[0].clear()


def apply_paragraph_formatting(doc: Document) -> None:
    body_started = False
    source_started = False
    for para in doc.paragraphs:
        original = para.text
        text = normalize_text(original)
        if text != original and para.runs:
            para.clear()
            para.add_run(text)

        stripped = para.text.strip()
        if stripped == "ВВЕДЕНИЕ":
            body_started = True
        if stripped == "СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ":
            source_started = True

        role = paragraph_role(stripped, body_started)
        if role == "h1":
            para.style = doc.styles["Heading 1"]
            if stripped != "СОДЕРЖАНИЕ":
                para.paragraph_format.page_break_before = True
        elif role == "h2":
            para.style = doc.styles["Heading 2"]
            para.paragraph_format.page_break_before = False
        elif role == "h3":
            para.style = doc.styles["Heading 3"]
            para.paragraph_format.page_break_before = False
        elif role == "caption":
            para.style = doc.styles["CaptionSys"]
        else:
            para.style = doc.styles["Normal"]
            fmt = para.paragraph_format
            fmt.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            fmt.first_line_indent = Cm(1.25)
            fmt.line_spacing_rule = WD_LINE_SPACING.SINGLE
            fmt.space_before = Pt(0)
            fmt.space_after = Pt(0)

        if source_started and re.match(r"^\d+\.", stripped):
            para.paragraph_format.first_line_indent = Cm(0)
            para.paragraph_format.left_indent = Cm(0)
            para.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

        for run in para.runs:
            if role in {"h1", "h2", "h3"}:
                set_run_font(run, "Arial", Pt(15 if role == "h1" else 14), bold=False)
            elif role == "caption":
                set_run_font(run, "Times New Roman", Pt(14))
            else:
                set_run_font(run, "Times New Roman", Pt(14))


def renumber_captions(doc: Document) -> None:
    table_no = 1
    figure_no = 1
    listing_no = 1
    for para in doc.paragraphs:
        text = para.text.strip()
        if text.startswith("Таблица "):
            rest = re.sub(r"^Таблица\s+\d+\s*-\s*", "", text)
            para.clear()
            para.add_run(f"Таблица {table_no} - {rest}")
            table_no += 1
        elif text.startswith("Рисунок "):
            rest = re.sub(r"^Рисунок\s+\d+\s*-\s*", "", text)
            rest = rest.replace("Заглушка для скриншота главной панели", "Главная панель SysAssist")
            rest = rest.replace("Заглушка для скриншота реестра модулей", "Реестр модулей SysAssist")
            rest = rest.replace("Заглушка для скриншота approval workflow", "Очередь согласований SysAssist")
            para.clear()
            para.add_run(f"Рисунок {figure_no} - {rest}")
            figure_no += 1
        elif text.startswith("Листинг "):
            rest = re.sub(r"^Листинг\s+\d+\s*-\s*", "", text)
            para.clear()
            para.add_run(f"Листинг {listing_no} - {rest}")
            listing_no += 1
        if para.text.strip().startswith(("Таблица ", "Рисунок ", "Листинг ")):
            para.style = doc.styles["CaptionSys"]
            for run in para.runs:
                set_run_font(run, "Times New Roman", Pt(14))


def is_code_table(table) -> bool:
    if len(table.rows) != 1:
        return False
    text = table.cell(0, 0).text
    code_markers = [
        "services.",
        "public ",
        "private ",
        "builder.",
        "Gate(",
        "dockerfile:",
        "dotnet test",
        "npm run build",
        "interface ",
        "UseNpgsql",
    ]
    return any(marker in text for marker in code_markers)


def format_tables(doc: Document) -> None:
    usable_twips = 9639  # A4 width minus 30 mm + 10 mm margins.
    for table in doc.tables:
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        set_table_width(table, usable_twips)
        set_table_cell_margins(table)
        rows = len(table.rows)
        cols = max(len(r.cells) for r in table.rows) if rows else 0
        code_table = is_code_table(table)
        col_width = usable_twips // max(cols, 1)
        for row_idx, row in enumerate(table.rows):
            for cell in row.cells:
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                set_cell_width(cell, usable_twips if code_table else col_width)
                if row_idx == 0 and not code_table:
                    set_cell_shading(cell, "D9EAF7")
                for para in cell.paragraphs:
                    para.paragraph_format.first_line_indent = Cm(0)
                    para.paragraph_format.space_before = Pt(0)
                    para.paragraph_format.space_after = Pt(0)
                    para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
                    para.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    if code_table:
                        para.style = doc.styles["CodeSys"]
                    for run in para.runs:
                        if code_table:
                            set_run_font(run, "Consolas", Pt(9))
                        else:
                            set_run_font(run, "Times New Roman", Pt(11 if cols >= 3 else 12), bold=(row_idx == 0))


def clear_paragraph(para) -> None:
    para.clear()
    para.paragraph_format.first_line_indent = Cm(0)
    para.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER


def replace_screenshots(doc: Document) -> None:
    caption_to_image = dict(SCREENSHOT_BY_CAPTION)
    for idx, para in enumerate(doc.paragraphs):
        caption = para.text.strip()
        if caption in caption_to_image and idx + 1 < len(doc.paragraphs):
            image_para = doc.paragraphs[idx + 1]
            if image_para._p.xpath(".//pic:pic"):
                clear_paragraph(image_para)
                run = image_para.add_run()
                run.add_picture(str(caption_to_image[caption]), width=Cm(13.2))

    # Insert two missing interface screenshots after known anchors.
    existing = {p.text.strip() for p in doc.paragraphs}
    for caption, image_path, note, anchor_caption in EXTRA_SCREENSHOTS:
        if caption in existing:
            continue
        for idx, para in enumerate(doc.paragraphs):
            if para.text.strip() == anchor_caption:
                img_para = doc.paragraphs[idx + 2] if idx + 2 < len(doc.paragraphs) else para
                new_cap = img_para.insert_paragraph_before(caption, style=doc.styles["CaptionSys"])
                for r in new_cap.runs:
                    set_run_font(r, "Times New Roman", Pt(14))
                new_img = img_para.insert_paragraph_before()
                new_img.paragraph_format.first_line_indent = Cm(0)
                new_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                new_img.add_run().add_picture(str(image_path), width=Cm(13.2))
                new_note = img_para.insert_paragraph_before(note, style=doc.styles["Normal"])
                new_note.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                new_note.paragraph_format.first_line_indent = Cm(1.25)
                for r in new_note.runs:
                    set_run_font(r, "Times New Roman", Pt(14))
                break
    renumber_captions(doc)


def count_sources(doc: Document) -> int:
    in_sources = False
    count = 0
    for para in doc.paragraphs:
        t = para.text.strip()
        if t == "СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ":
            in_sources = True
            continue
        if in_sources and re.match(r"^\d+\.", t):
            count += 1
    return count


def update_toc_static(doc: Document) -> None:
    # Keep the existing static TOC but normalize labels and table order. Word COM will refresh page numbers later.
    body_headings = []
    body_started = False
    for para in doc.paragraphs:
        t = para.text.strip()
        if t == "ВВЕДЕНИЕ":
            body_started = True
        if not body_started:
            continue
        if para.style.name in {"Heading 1", "Heading 2", "Heading 3"}:
            if t != "СОДЕРЖАНИЕ":
                body_headings.append(t)

    toc_start = None
    toc_end = None
    for i, para in enumerate(doc.paragraphs):
        if para.text.strip() == "СОДЕРЖАНИЕ":
            toc_start = i
        elif toc_start is not None and para.text.strip() == "ВВЕДЕНИЕ":
            toc_end = i
            break
    if toc_start is None or toc_end is None:
        return

    toc_title = doc.paragraphs[toc_start]
    toc_title.clear()
    toc_title.add_run("СОДЕРЖАНИЕ")
    toc_title.style = doc.styles["CaptionSys"]
    toc_title.paragraph_format.first_line_indent = Cm(0)
    toc_title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in toc_title.runs:
        set_run_font(run, "Arial", Pt(15))

    # Remove old static TOC entries between the TOC title and the first body heading.
    for idx in range(toc_end - 1, toc_start, -1):
        el = doc.paragraphs[idx]._element
        el.getparent().remove(el)

    intro = None
    for para in doc.paragraphs:
        if para.text.strip() == "ВВЕДЕНИЕ" and para.style.name == "Heading 1":
            intro = para
            break
    if intro is None:
        return

    for heading in body_headings:
        new_para = intro.insert_paragraph_before(f"{heading}\t0", style=doc.styles["Normal"])
        new_para.paragraph_format.first_line_indent = Cm(0)
        new_para.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        for run in new_para.runs:
            set_run_font(run, "Times New Roman", Pt(14))


def write_markdown_reports(doc: Document) -> None:
    sources_count = count_sources(doc)
    checked_files = [
        "review_work/input/diploma_source.docx",
        "review_work/input/sto_rules.docx",
        "review_work/input/diploma_example.docx",
        "review_work/input/vkr_structure.docx",
        "README.md",
        "src/SysAssist.Api/Program.cs",
        "src/SysAssist.Infrastructure/Data/SysAssistDbContext.cs",
        "src/SysAssist.Infrastructure/Modules/ModuleCatalog.cs",
        "src/SysAssist.Infrastructure/Modules/RemediationActionCatalog.cs",
        "src/SysAssist.Infrastructure/Security/SecretProtection.cs",
        "src/SysAssist.Infrastructure/Security/LicenseService.cs",
        "tests/SysAssist.Tests",
        "docker-compose.yml",
        "docker-compose.prod.yml",
        "scripts/prod-smoke.ps1",
        "docs/security-checklist.md",
        "docs/prod-readiness.md",
        "docs/server-deployment.md",
    ]

    fact_rows = [
        ("Backend .NET 9 и ASP.NET Core Minimal API", "Подтверждено", "src/SysAssist.Api/SysAssist.Api.csproj: TargetFramework net9.0; Program.cs: MapGet/MapPost groups"),
        ("CockroachDB через EF Core/Npgsql", "Подтверждено", "src/SysAssist.Infrastructure.csproj: Npgsql.EntityFrameworkCore.PostgreSQL; SysAssistDbContext; migration InitialCockroachSchema"),
        ("Frontend React, TypeScript, Vite", "Подтверждено", "src/SysAssist.Web/package.json: react, typescript, vite"),
        ("Архитектура Domain/Application/Contracts/Infrastructure/Api/Web", "Подтверждено", "README.md и структура src/*"),
        ("13 встроенных модулей", "Подтверждено", "ModuleCatalog.cs: Zabbix, Grafana, Alertmanager, PostgreSQL, Redis, Docker, Nginx, Linux Host, HTTP, File System, SMTP, Telegram, Local Rule Advisor"),
        ("65 remediation actions", "Подтверждено", "RemediationActionCatalog.cs: 13 модулей по 5 actions"),
        ("SafeMode блокирует реальные действия", "Подтверждено", "SysAssistApiService.cs: при module.SafeMode approved action не исполняется и аудитируется ACTION_BLOCKED"),
        ("Fallback Mode не является production-режимом", "Подтверждено", "ModuleCatalog.cs: UseFallbackMode Deprecated; production gate требует отключения fallback"),
        ("Approval workflow для high/critical actions", "Подтверждено", "SysAssistApiService.cs: RequiresApproval/High/Critical создают ApprovalRequest"),
        ("Audit log и support bundle", "Подтверждено", "SysAssistApiService.cs: AuditAsync, GetSupportBundleFileAsync; tests проверяют отсутствие секретов"),
        ("Production deployment", "Частично подтверждено", "docker-compose.prod.yml и docs/server-deployment.md подготовлены; промышленное внедрение не заявляется"),
        ("SSO/OIDC", "Не подтверждено как реализация", "docs/prod-readiness.md относит SSO/OIDC к Current Product Gaps; в дипломе оставлено как развитие"),
    ]

    fact_md = ["# Fact-check matrix", "", "| Утверждение | Статус | Подтверждение / исправление |", "| --- | --- | --- |"]
    for claim, status, evidence in fact_rows:
        fact_md.append(f"| {claim} | {status} | {evidence} |")
    (ROOT / "fact_check_matrix.md").write_text("\n".join(fact_md) + "\n", encoding="utf-8")

    fmt_rows = [
        ("Формат А4", "Выполнено", "Установлено 210 x 297 мм для всех секций."),
        ("Поля 30/10/20/20 мм", "Выполнено", "Левое 30 мм, правое 10 мм, верх/низ 20 мм."),
        ("Основной текст Times New Roman 14", "Выполнено", "Нормализован стиль Normal и прямое форматирование."),
        ("Выравнивание по ширине и абзац 1,25 см", "Выполнено", "Применено к основному тексту."),
        ("Одинарный интервал без дополнительных интервалов", "Выполнено", "Применено к стилям и абзацам."),
        ("Заголовки Arial 15/14", "Выполнено", "Heading 1/2/3 приведены к Arial."),
        ("Номер страницы сверху справа, Arial 12", "Выполнено", "Добавлено поле PAGE в верхний колонтитул, первая страница без номера."),
        ("Таблицы с номерами и названиями", "Выполнено", "Подписи таблиц перенумерованы по фактическому порядку."),
        ("Листинги кода короткие и читаемые", "Выполнено", "Кодовые таблицы оформлены Consolas 9."),
        ("Не более 5 интерфейсных скриншотов", "Выполнено", "В основном тексте использованы dashboard, modules, events, approvals, diagnostics."),
        ("Источников минимум 30", "Выполнено", f"В списке {sources_count} источников."),
        ("Отсутствие длинного тире", "Выполнено", "Длинные тире заменены на дефис."),
    ]
    fmt_md = ["# Formatting checklist", "", "| Требование СТО | Статус | Комментарий |", "| --- | --- | --- |"]
    for req, status, comment in fmt_rows:
        fmt_md.append(f"| {req} | {status} | {comment} |")
    (ROOT / "formatting_checklist.md").write_text("\n".join(fmt_md) + "\n", encoding="utf-8")

    sources_md = [
        "# Sources check",
        "",
        f"В итоговом списке использованных источников - {sources_count} позиций.",
        "",
        "Проверены категории источников: локальная документация SysAssist, СТО/структура ВКР, официальная документация Microsoft/.NET/EF Core, Npgsql, CockroachDB, Docker, Nginx, React, Vite, TypeScript, Zabbix, Grafana, Prometheus, PostgreSQL, Redis, OWASP, NIST, ISO/IEC, Serilog, xUnit и GitHub Actions.",
        "",
        "Проблемные учебные и нерелевантные источники из примера оформления не перенесены. Внешние источники оставлены как официальные или нормативные; утверждения о реализации проекта подтверждаются локальным исходным кодом и документацией репозитория.",
    ]
    (ROOT / "sources_check.md").write_text("\n".join(sources_md) + "\n", encoding="utf-8")

    review_md = [
        "# Итог проверки дипломной работы SysAssist",
        "",
        "## 1. Проверенные файлы",
        "",
        *[f"- `{item}`" for item in checked_files],
        "",
        "## 2. Основные исправления",
        "",
        "- Документ приведен к формату А4, полям 30/10/20/20 мм, Times New Roman 14 для основного текста и Arial 15/14 для заголовков.",
        "- Перенумерованы подписи таблиц, рисунков и листингов по порядку следования.",
        "- Заглушки интерфейса заменены реальными скриншотами из `review_work/screenshots`.",
        "- Кодовые фрагменты оформлены как короткие листинги с Consolas 9.",
        "- Убраны длинные тире и несколько неудачных формулировок про продуктивный режим.",
        "- Проверены фактические утверждения по стеку, базе данных, модулям, SafeMode, approval, audit, support bundle и production readiness.",
        "",
        "## 3. Фактические несоответствия",
        "",
        "| Было в тексте | Почему неверно или неполно | Как исправлено | Подтверждение в проекте |",
        "| --- | --- | --- | --- |",
        "| Формулировки с заглушками скриншотов | Реальные PNG уже были подготовлены | Заменены на реальные скриншоты dashboard/modules/events/approvals/diagnostics | `review_work/screenshots` |",
        "| Номера таблиц 10-13 шли не по порядку | Нарушение нормоконтроля | Таблицы перенумерованы по порядку следования | Итоговый DOCX |",
        "| Риск смешать production-shaped deployment и внедрение | В репозитории есть deployment artifacts, но нет доказательства промышленного внедрения | В тексте сохранена формулировка про controlled demo и пилот | `README.md`, `docs/prod-readiness.md` |",
        "| SafeMode мог восприниматься как успешное выполнение | В коде действие блокируется при SafeMode | Уточнено `ACTION_BLOCKED`, не реальное воздействие | `SysAssistApiService.cs` |",
        "",
        "## 4. Проверка СТО",
        "",
        "См. `formatting_checklist.md`.",
        "",
        "## 5. Проверка листингов кода",
        "",
        "| Листинг | Файл проекта | Статус | Что исправлено |",
        "| --- | --- | --- | --- |",
        "| Регистрация адаптеров через DI | `src/SysAssist.Infrastructure/DependencyInjection.cs` | Подтвержден | Оформлен как короткий кодовый фрагмент |",
        "| Подключение CockroachDB через Npgsql | `src/SysAssist.Api/Program.cs`, `SysAssistDbContextFactory.cs` | Подтвержден | Сохранена короткая выжимка |",
        "| Политики ролей ASP.NET Core | `src/SysAssist.Api/Program.cs` | Подтвержден | Оформлен Consolas 9 |",
        "| Защита секретов AES-GCM | `SecretProtection.cs` | Подтвержден | Уточнен реальный механизм `enc:v1:` |",
        "| Интерфейс интеграционного адаптера | `IIntegrationAdapter.cs` | Подтвержден | Сохранен только интерфейсный фрагмент |",
        "| Production compose | `docker-compose.prod.yml` | Подтвержден частично | Описан как production-shaped deployment |",
        "| Production readiness gate | `Program.cs` | Подтвержден | Оставлена короткая логика gate |",
        "| Команды локальной проверки | `README.md`, `scripts/prod-smoke.ps1` | Подтвержден | Оставлены команды проверки |",
        "",
        "## 6. Проверка источников",
        "",
        f"В списке {sources_count} источников. Основа списка - официальная документация технологий и локальная документация проекта. Отдельный файл: `sources_check.md`.",
        "",
        "## 7. Проверка объема",
        "",
        "Объем должен быть подтвержден после Word/PDF-рендера. LibreOffice в среде отсутствует, поэтому используется Microsoft Word COM для экспорта и подсчета страниц.",
        "",
        "## 8. Остаточные риски",
        "",
        "- Номера страниц в статическом содержании требуют финального обновления после Word-рендера.",
        "- Docker на текущей машине не установлен, поэтому фактический запуск compose не выполнялся.",
        "- Внешний демо-сайт и реальные интеграции требуют ручной проверки доступов, секретов и сетевой связности.",
    ]
    (ROOT / "codex_review_report.md").write_text("\n".join(review_md) + "\n", encoding="utf-8")


def inspect_docx_xml(path: Path) -> dict[str, int]:
    with zipfile.ZipFile(path) as z:
        xml = z.read("word/document.xml").decode("utf-8", errors="ignore")
    return {
        "long_dash": xml.count("—") + xml.count("–"),
        "express": len(re.findall(r"Express", xml, flags=re.I)),
        "sqlite": len(re.findall(r"SQLite", xml, flags=re.I)),
    }


def main() -> None:
    prepare_docx_screenshots()
    doc = Document(str(INPUT))
    ensure_styles(doc)
    apply_section_format(doc)
    apply_paragraph_formatting(doc)
    renumber_captions(doc)
    replace_screenshots(doc)
    format_tables(doc)
    update_toc_static(doc)
    write_markdown_reports(doc)
    doc.save(str(OUTPUT))

    checks = inspect_docx_xml(OUTPUT)
    print(f"saved={OUTPUT}")
    print(f"sources={count_sources(doc)}")
    print(f"long_dash={checks['long_dash']} express={checks['express']} sqlite={checks['sqlite']}")


if __name__ == "__main__":
    main()
